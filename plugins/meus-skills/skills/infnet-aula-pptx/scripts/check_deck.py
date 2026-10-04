#!/usr/bin/env python3
"""Confere se um .pptx gerado obedece às regras de formatação da skill.

    python check_deck.py Arquivo.pptx

Erros (código de saída 1) são violações de regra. Avisos pedem um olhar humano:
a capitalização de subtítulos e o equilíbrio entre listas e imagens são heurísticas.
"""
import sys

from pptx import Presentation
from pptx.util import Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

POR_TEMA = 10
BOTTOM = int(5.12 * 914400)
MINUSCULAS = {"a", "o", "as", "os", "um", "uma", "de", "do", "da", "dos", "das", "e", "ou", "em", "no", "na",
              "nos", "nas", "com", "sem", "por", "para", "pelo", "pela", "ao", "à", "aos", "às", "que", "como",
              "entre", "sobre", "sob", "antes", "após", "até", "se"}


def run_ok(r, size):
    try:
        color = str(r.font.color.rgb)
    except Exception:
        color = None
    return r.font.name == "Calibri" and r.font.size == Pt(size) and color == "000000"


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    prs = Presentation(sys.argv[1])
    errs, warns = [], []
    n = len(prs.slides)
    if n == 0 or n % POR_TEMA:
        errs.append(f"{n} slides: o total precisa ser múltiplo de {POR_TEMA} (10 por tema)")
    lista = imagem = 0
    titulos = []
    for i, s in enumerate(prs.slides, 1):
        if s.slide_layout.name != "OBJECT":
            errs.append(f"slide {i}: layout '{s.slide_layout.name}' (esperado OBJECT)")
        title = s.shapes.title
        body = next((p for p in s.placeholders if p.placeholder_format.idx == 1), None)
        if title is None or body is None:
            errs.append(f"slide {i}: falta título ou corpo")
            continue
        titulos.append(title.text_frame.text)
        if len(title.text_frame.paragraphs) != 1 or not title.text_frame.text.strip():
            errs.append(f"slide {i}: título vazio ou com mais de um parágrafo")
        for r in title.text_frame.paragraphs[0].runs:
            if not run_ok(r, 30):
                errs.append(f"slide {i}: título fora do padrão Calibri 30 preta")
        paras = body.text_frame.paragraphs
        sub = paras[0]
        if "buNone" not in sub._p.xml or not sub.runs or not all(r.font.bold for r in sub.runs):
            errs.append(f"slide {i}: subtítulo precisa ser negrito e sem marcador")
        words = sub.text.replace(":", " ").split()
        for k, w in enumerate(words):
            base = w.strip("()“”\"'.,;?!")
            if not base or base[0].isdigit():
                continue
            if k > 0 and base.lower() in MINUSCULAS:
                if base != base.lower():
                    warns.append(f"slide {i}: subtítulo '{sub.text}': '{base}' deveria estar em minúsculas?")
            elif not base[0].isupper():
                warns.append(f"slide {i}: subtítulo '{sub.text}': '{base}' deveria iniciar em maiúscula?")
            elif base[1:] != base[1:].lower():
                warns.append(f"slide {i}: subtítulo '{sub.text}': '{base}' tem maiúsculas no meio (sigla?)")
        tem_lista = False
        for k, p in enumerate(paras):
            for r in p.runs:
                if not run_ok(r, 18):
                    errs.append(f"slide {i}: texto fora do padrão Calibri 18 preta: '{r.text[:30]}'")
                if k > 0 and r.font.bold and r.text.strip() == p.text.strip() and "buNone" in p._p.xml and len(paras) > 2:
                    warns.append(f"slide {i}: parágrafo inteiro em negrito no corpo: '{p.text[:30]}' (segundo subtítulo?)")
            if "buChar" in p._p.xml or "buAutoNum" in p._p.xml:
                tem_lista = True
        lista += tem_lista
        pics = [sh for sh in s.shapes if sh.shape_type == 13]
        imagem += bool(pics)
        for sh in pics:
            if sh.top + sh.height > BOTTOM or sh.left < 0 or sh.left + sh.width > prs.slide_width:
                errs.append(f"slide {i}: imagem fora da área útil")
            if not sh._element.nvPicPr.cNvPr.get("descr"):
                warns.append(f"slide {i}: imagem sem texto alternativo")
        pos = (i - 1) % POR_TEMA + 1
        if pos == POR_TEMA and "exerc" not in sub.text.lower():
            errs.append(f"slide {i}: o último slide do tema deve ser de exercícios de fixação (subtítulo: '{sub.text}')")
        if pos == POR_TEMA - 1 and "exemplo pr" not in sub.text.lower():
            warns.append(f"slide {i}: o slide 9 do tema deveria ser o exemplo prático (subtítulo: '{sub.text}')")
    for t in range(0, len(titulos) - len(titulos) % POR_TEMA, POR_TEMA):
        if len(set(titulos[t:t + POR_TEMA])) != 1:
            errs.append(f"tema {t // POR_TEMA + 1}: os 10 slides precisam do mesmo título")
    if n:
        if lista == n:
            errs.append("todos os slides usam lista; marcadores devem aparecer só quando necessários")
        elif lista > 0.6 * n:
            warns.append(f"{lista} de {n} slides usam lista; reduza para perto da metade")
        if imagem < 0.3 * n:
            warns.append(f"só {imagem} de {n} slides têm imagem; o alvo é cerca de 5 por tema")
    print(f"{n} slides | {n // POR_TEMA} tema(s) | {lista} com lista | {imagem} com imagem")
    for w in warns:
        print("AVISO:", w)
    for e in errs:
        print("ERRO:", e)
    print("Resultado:", "REPROVADO" if errs else "aprovado" + (" com avisos" if warns else ""))
    return 1 if errs else 0


if __name__ == "__main__":
    sys.exit(main())
