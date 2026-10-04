#!/usr/bin/env python3
"""Monta o .pptx da aula a partir de uma especificação JSON, sobre assets/exemplo.pptx.

    python build_deck.py spec.json [--out-dir DIR]

O arquivo é gravado na raiz da skill (pasta acima de scripts/). Se essa pasta não
aceitar escrita, o script grava no diretório atual e avisa.

Formato da especificação (caminhos de imagem relativos ao arquivo spec):

    {
      "arquivo": "Nome_Do_Arquivo.pptx",
      "temas": [
        {
          "titulo": "Título Curto do Tema",
          "slides": [
            {
              "subtitulo": "Tema 1: Visão Geral",
              "corpo": [["p", "Parágrafo com **negrito** e *itálico*."],
                        ["b", "item com marcador;"],
                        ["n", "item numerado;"]],
              "imagem": {"arquivo": "img/f01.png", "posicao": "abaixo", "alt": "Descrição."}
            }
          ]
        }
      ]
    }

Cada tema precisa de exatamente 10 slides. "posicao" é "abaixo" (imagem sob o texto)
ou "lado" (texto à esquerda, imagem à direita). "imagem" é opcional.

O script mede o texto com as métricas da Calibri e recusa (código de saída 2) o que
não cabe: título ou subtítulo com mais de uma linha e corpo maior que a área útil.
"""
import argparse
import json
import os
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SKILL_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_ROOT / "assets" / "exemplo.pptx"
LAYOUT = "OBJECT"
FONT = "Calibri"
BLACK = RGBColor(0, 0, 0)
EMU_IN = 914400
SLIDES_POR_TEMA = 10

TITLE_PT, TEXT_PT = 30, 18
LINE_PT = 22.0                         # altura de linha da Calibri 18 pt
INS_L, INS_T = 68575, 34275            # margens internas das caixas de texto do template
BOTTOM_LIMIT = int(5.08 * EMU_IN)      # limite inferior útil do cartão branco do slide
SIDE_TEXT_W = int(4.25 * EMU_IN)       # largura do texto quando a imagem fica ao lado
TITLE_MAX_IN = 7.9                     # além disso o título encosta no logotipo
MAX_UPSCALE = 1.12

# ----------------------------------------------------------------- medição de texto
_FONT_FILES = {
    (0, 0): ["Carlito-Regular.ttf", "calibri.ttf", "Calibri.ttf"],
    (1, 0): ["Carlito-Bold.ttf", "calibrib.ttf", "Calibri Bold.ttf"],
    (0, 1): ["Carlito-Italic.ttf", "calibrii.ttf", "Calibri Italic.ttf"],
    (1, 1): ["Carlito-BoldItalic.ttf", "calibriz.ttf", "Calibri Bold Italic.ttf"],
}
_FONT_DIRS = [
    "/usr/share/fonts/truetype/crosextra",
    os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
    "/Library/Fonts", "/Library/Fonts/Microsoft", os.path.expanduser("~/Library/Fonts"),
    "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts",
    "/usr/share/fonts/carlito", "/usr/share/fonts/google-carlito-fonts", "/usr/share/fonts/truetype/carlito",
]
_cache = {}
METRICAS_REAIS = True


def _font(bold, italic, size):
    key = (int(bold), int(italic), size)
    if key not in _cache:
        found = None
        for d in _FONT_DIRS:
            for name in _FONT_FILES[key[:2]]:
                p = os.path.join(d, name)
                if os.path.isfile(p):
                    found = p
                    break
            if found:
                break
        _cache[key] = ImageFont.truetype(found, size * 10) if found else None
    return _cache[key]


def text_width_pt(text, bold=False, italic=False, size=TEXT_PT):
    f = _font(bold, italic, size)
    if f is None:                      # sem Calibri/Carlito: estimativa conservadora
        global METRICAS_REAIS
        METRICAS_REAIS = False
        return len(text) * size * (0.53 if bold else 0.51)
    return f.getlength(text) / 10


def parse(text):
    """'**negrito**' e '*itálico*' -> [(texto, negrito, itálico)]"""
    out, bold, ital = [], False, False
    for tok in re.split(r"(\*\*|\*)", text):
        if tok == "**":
            bold = not bold
        elif tok == "*":
            ital = not ital
        elif tok:
            out.append((tok, bold, ital))
    return out


def count_lines(runs, width_pt, size=TEXT_PT):
    lines, cur = 1, 0.0
    for t, b, i in runs:
        for m in re.finditer(r"\S+\s*", t):
            w = m.group(0)
            if cur + text_width_pt(w.rstrip(), b, i, size) > width_pt and cur > 0:
                lines += 1
                cur = 0.0
            cur += text_width_pt(w, b, i, size)
    return lines


def spacing(kind, prev):
    """Espaço antes do parágrafo, em pontos."""
    if prev == "sub":
        return 10
    if kind == "p":
        return 7
    return 4 if prev not in ("b", "n") else 3


# ----------------------------------------------------------------- escrita no pptx
def style_run(r, size, bold=False, italic=False):
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = BLACK
    rPr = r._r.get_or_add_rPr()
    rPr.set("lang", "pt-BR")
    for tag in ("a:ea", "a:cs", "a:sym"):
        etree.SubElement(rPr, qn(tag)).set("typeface", FONT)


def set_ppr(p, kind, spc_bef):
    pPr = p._p.get_or_add_pPr()
    for k in list(pPr.attrib):
        del pPr.attrib[k]
    for ch in list(pPr):
        pPr.remove(ch)
    pPr.set("lvl", "0")
    pPr.set("algn", "l")
    pPr.set("rtl", "0")
    if kind in ("b", "n"):
        pPr.set("marL", "457200")
        pPr.set("indent", "-342900")
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
    sb = etree.SubElement(pPr, qn("a:spcBef"))
    etree.SubElement(sb, qn("a:spcPts")).set("val", str(int(spc_bef * 100)))
    sa = etree.SubElement(pPr, qn("a:spcAft"))
    etree.SubElement(sa, qn("a:spcPts")).set("val", "0")
    if kind == "b":
        etree.SubElement(pPr, qn("a:buSzPts")).set("val", "1800")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "●")
    elif kind == "n":
        etree.SubElement(pPr, qn("a:buSzPts")).set("val", "1800")
        etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
    else:                              # subtítulo e parágrafos nunca levam marcador
        etree.SubElement(pPr, qn("a:buNone"))


def end_rpr(p, size, bold=False):
    e = etree.SubElement(p._p, qn("a:endParaRPr"))
    e.set("sz", str(size * 100))
    if bold:
        e.set("b", "1")


def image_inches(path):
    with Image.open(path) as im:
        w, h = im.size
        dpi = im.info.get("dpi", (220, 220))[0] or 220
    return w / dpi, h / dpi


# ----------------------------------------------------------------- validação da spec
def validate_spec(spec, base):
    errs = []
    temas = spec.get("temas")
    if not isinstance(temas, list) or not temas:
        return ["a especificação precisa de uma lista 'temas' não vazia"]
    for ti, tema in enumerate(temas, 1):
        tit = tema.get("titulo", "")
        if not tit:
            errs.append(f"tema {ti}: falta 'titulo'")
        elif text_width_pt(tit, size=TITLE_PT) / 72 > TITLE_MAX_IN:
            errs.append(f"tema {ti}: título '{tit}' tem {text_width_pt(tit, size=TITLE_PT) / 72:.1f} pol "
                        f"em 30 pt (máximo {TITLE_MAX_IN}); encurte-o")
        slides = tema.get("slides", [])
        if len(slides) != SLIDES_POR_TEMA:
            errs.append(f"tema {ti}: {len(slides)} slides; cada tema precisa de exatamente {SLIDES_POR_TEMA}")
        for si, sl in enumerate(slides, 1):
            where = f"tema {ti}, slide {si}"
            if not sl.get("subtitulo"):
                errs.append(f"{where}: falta 'subtitulo'")
            if not sl.get("corpo"):
                errs.append(f"{where}: falta 'corpo'")
            for item in sl.get("corpo", []):
                if not (isinstance(item, list) and len(item) == 2 and item[0] in ("p", "b", "n")):
                    errs.append(f"{where}: item de corpo inválido {item!r} (use ['p'|'b'|'n', 'texto'])")
            img = sl.get("imagem")
            if img:
                if img.get("posicao") not in ("abaixo", "lado"):
                    errs.append(f"{where}: imagem.posicao deve ser 'abaixo' ou 'lado'")
                if not (base / img.get("arquivo", "")).is_file():
                    errs.append(f"{where}: imagem não encontrada: {img.get('arquivo')}")
    return errs


# ----------------------------------------------------------------- montagem
def build(spec, base, out_path):
    prs = Presentation(str(TEMPLATE))
    ids = prs.slides._sldIdLst
    for sldId in list(ids):            # remove os slides de exemplo; ficam só layout e tema
        prs.part.drop_rel(sldId.rId)
        ids.remove(sldId)
    layout = next(l for l in prs.slide_layouts if l.name == LAYOUT)
    body_ph = next(p for p in layout.placeholders if p.placeholder_format.idx == 1)
    BX, BY, BW, BH = body_ph.left, body_ph.top, body_ph.width, body_ph.height
    avail_pt = (BH - 2 * INS_T) / 12700

    errs, warns, rows, n = [], [], [], 0
    for ti, tema in enumerate(spec["temas"], 1):
        for si, sl in enumerate(tema["slides"], 1):
            n += 1
            where = f"slide {n} (tema {ti}, slide {si})"
            s = prs.slides.add_slide(layout)
            p = s.shapes.title.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = tema["titulo"]
            style_run(r, TITLE_PT)
            end_rpr(p, TITLE_PT)

            img = sl.get("imagem")
            side = bool(img) and img["posicao"] == "lado"
            width = SIDE_TEXT_W if side else BW
            inner = (width - 2 * INS_L) / 12700
            body = s.placeholders[1]
            body.left, body.top, body.width, body.height = BX, BY, width, BH
            tf = body.text_frame
            tf.word_wrap = True

            p = tf.paragraphs[0]
            set_ppr(p, "sub", 0)
            r = p.add_run()
            r.text = sl["subtitulo"]
            style_run(r, TEXT_PT, bold=True)
            end_rpr(p, TEXT_PT, bold=True)
            if count_lines([(sl["subtitulo"], True, False)], inner) > 1:
                errs.append(f"{where}: subtítulo não cabe em uma linha"
                            + (" na coluna estreita; encurte-o ou use posicao 'abaixo'" if side else "; encurte-o"))

            h_pt, prev, detail = LINE_PT, "sub", []
            for kind, text in sl["corpo"]:
                p = tf.add_paragraph()
                set_ppr(p, kind, spacing(kind, prev))
                runs = parse(text)
                for t, b, i in runs:
                    r = p.add_run()
                    r.text = t
                    style_run(r, TEXT_PT, bold=b, italic=i)
                end_rpr(p, TEXT_PT)
                lines = count_lines(runs, (inner - (36 if kind in ("b", "n") else 0)) * 0.985)
                h_pt += spacing(kind, prev) + lines * LINE_PT
                detail.append(lines)
                prev = kind

            note = ""
            if not img and h_pt > avail_pt:
                errs.append(f"{where}: texto ocupa {h_pt:.0f} pt e a área tem {avail_pt:.0f} pt; "
                            f"corte cerca de {int((h_pt - avail_pt) / LINE_PT) + 1} linha(s)")
            if img:
                path = base / img["arquivo"]
                w_in, h_in = image_inches(path)
                w_emu, h_emu = int(w_in * EMU_IN), int(h_in * EMU_IN)
                if side:
                    if h_pt > avail_pt:
                        errs.append(f"{where}: texto da coluna ocupa {h_pt:.0f} pt e a área tem {avail_pt:.0f} pt")
                    x0 = BX + SIDE_TEXT_W + int(0.12 * EMU_IN)
                    room_w = BX + BW - x0
                    room_h = BOTTOM_LIMIT - BY - int(0.1 * EMU_IN)
                    sc = min(1.0, room_w / w_emu, room_h / h_emu)
                    w_emu, h_emu = int(w_emu * sc), int(h_emu * sc)
                    left = x0 + (room_w - w_emu) // 2
                    top = BY + int(0.1 * EMU_IN) + (room_h - h_emu) // 2
                else:
                    top0 = BY + INS_T + int(h_pt * 12700) + int(0.16 * EMU_IN)
                    room_h = BOTTOM_LIMIT - top0
                    if room_h < int(0.9 * EMU_IN):
                        errs.append(f"{where}: sobra {max(room_h, 0) / EMU_IN:.1f} pol para a imagem; "
                                    f"reduza o texto (até umas 4 linhas) ou use posicao 'lado'")
                        room_h = int(0.9 * EMU_IN)
                    sc = min(MAX_UPSCALE, (BW - 2 * INS_L) / w_emu, room_h / h_emu)
                    w_emu, h_emu = int(w_emu * sc), int(h_emu * sc)
                    left = BX + (BW - w_emu) // 2
                    top = top0 + max(0, (room_h - h_emu) // 2)
                note = f"imagem x{sc:.2f}"
                if sc < 0.8:
                    warns.append(f"{where}: imagem reduzida a {sc:.0%}; o texto dentro dela fica pequeno. "
                                 f"Redesenhe a figura no tamanho indicado ou encurte o texto do slide")
                pic = s.shapes.add_picture(str(path), left, top, w_emu, h_emu)
                pic.name = "Diagrama " + Path(img["arquivo"]).stem
                if img.get("alt"):
                    pic._element.nvPicPr.cNvPr.set("descr", img["alt"])
                else:
                    warns.append(f"{where}: imagem sem 'alt' (descrição para leitores de tela)")
            rows.append((n, round(h_pt), round(avail_pt), detail, note, sl["subtitulo"]))

    prs.core_properties.title = spec.get("titulo_apresentacao") or " / ".join(t["titulo"] for t in spec["temas"])
    prs.save(str(out_path))
    return rows, errs, warns


def resolve_out(spec, out_dir):
    name = spec.get("arquivo") or "Aula.pptx"
    if not name.lower().endswith(".pptx"):
        name += ".pptx"
    name = Path(name).name             # sempre só o nome: o destino é decidido aqui
    if out_dir:
        d = Path(out_dir)
        d.mkdir(parents=True, exist_ok=True)
        return d / name, None
    if os.access(SKILL_ROOT, os.W_OK):
        return SKILL_ROOT / name, None
    return Path.cwd() / name, (f"AVISO: a raiz da skill ({SKILL_ROOT}) não aceita escrita; "
                               f"o arquivo foi gravado em {Path.cwd()}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("spec", help="arquivo JSON com a especificação dos slides")
    ap.add_argument("--out-dir", help="grava em outro diretório em vez da raiz da skill")
    a = ap.parse_args()

    spec_path = Path(a.spec).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    errs = validate_spec(spec, spec_path.parent)
    if errs:
        print("ESPECIFICAÇÃO INVÁLIDA:")
        for e in errs:
            print("  -", e)
        return 2

    out_path, aviso = resolve_out(spec, a.out_dir)
    rows, errs, warns = build(spec, spec_path.parent, out_path)
    for n, h, avail, detail, note, sub in rows:
        print(f"{n:3d}  {h:3d}/{avail} pt  linhas {detail}  {note:13s} {sub}")
    if not METRICAS_REAIS:
        print("AVISO: Calibri/Carlito não encontrada; medidas de texto são estimativas. Confira a renderização.")
    for w in warns:
        print("AVISO:", w)
    if aviso:
        print(aviso)
    total = len(rows)
    print(f"\n{total} slides ({len(spec['temas'])} tema(s) x {SLIDES_POR_TEMA}) gravados em: {out_path}")
    if errs:
        print("\nERROS DE ENCAIXE (corrija a especificação e gere de novo):")
        for e in errs:
            print("  -", e)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
